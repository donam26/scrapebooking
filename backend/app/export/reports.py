# ruff: noqa: E501 — mẫu HTML in được giữ nguyên dòng cho dễ đọc
"""Báo cáo đúng mẫu khách sạn dùng (roadmap 7.4, BC N13):

- **Excel "rate shop"**: khách sạn × đêm (giá niêm yết, trạng thái 5 mức, KM, số đêm tối thiểu),
  trang compset theo đêm (n/N, trung vị, chỉ số giá niêm yết, vị trí giá) và trang thay đổi.
- **Báo cáo tháng cho chủ đầu tư** (HTML in được thành PDF từ trình duyệt): KPI thật (PMS), vị trí
  giá so compset, gợi ý đã áp dụng → kết quả, sự kiện nổi bật của đối thủ.
"""

import io
from collections.abc import Iterable, Sequence
from dataclasses import dataclass
from datetime import date
from decimal import Decimal
from html import escape
from typing import Any

STATE_LABEL = {
    "vi": {
        "available": "Còn bán",
        "sold_out": "Hết phòng",
        "restricted": "Hạn chế",
        "no_price": "Không có giá",
        "error": "Lỗi đọc",
        None: "",
    },
    "en": {
        "available": "Available",
        "sold_out": "Sold out",
        "restricted": "Restricted",
        "no_price": "No price",
        "error": "Error",
        None: "",
    },
}

HEAD = {
    "vi": {
        "sheet_rates": "Giá đối thủ",
        "sheet_compset": "Compset",
        "sheet_changes": "Thay đổi",
        "hotel": "Khách sạn",
        "role": "Vai trò",
        "night": "Đêm",
        "state": "Trạng thái",
        "price": "Giá niêm yết",
        "rooms_left": "Phòng còn (kênh báo)",
        "min_stay": "Tối thiểu (đêm)",
        "promos": "Khuyến mãi",
        "observed": "Quan sát lúc",
        "n_priced": "Đối thủ có giá (n)",
        "n_total": "Đối thủ (N)",
        "sold_out": "Hết phòng",
        "restricted": "Hạn chế",
        "median": "Trung vị đối thủ",
        "own": "Giá của bạn",
        "index": "Chỉ số giá niêm yết",
        "rank": "Vị trí giá (1 = rẻ nhất)",
        "sample": "Cỡ mẫu",
        "event": "Thay đổi",
        "change": "Từ → tới",
        "reason": "Lý do",
        "channel": "Kênh",
    },
    "en": {
        "sheet_rates": "Competitor rates",
        "sheet_compset": "Compset",
        "sheet_changes": "Changes",
        "hotel": "Hotel",
        "role": "Role",
        "night": "Night",
        "state": "State",
        "price": "Advertised rate",
        "rooms_left": "Rooms left (channel)",
        "min_stay": "Min stay (nights)",
        "promos": "Promotions",
        "observed": "Observed at",
        "n_priced": "Competitors priced (n)",
        "n_total": "Competitors (N)",
        "sold_out": "Sold out",
        "restricted": "Restricted",
        "median": "Competitor median",
        "own": "Your rate",
        "index": "Advertised rate index",
        "rank": "Rate position (1 = cheapest)",
        "sample": "Sample",
        "event": "Change",
        "change": "From → to",
        "reason": "Reason",
        "channel": "Channel",
    },
}


def _num(v: Any) -> float | None:
    return float(v) if v is not None else None


def rate_shop_xlsx(overview: Any, events: Sequence[Any], locale: str = "vi") -> bytes:
    """Workbook từ OverviewOut (khách sạn × đêm + compset) và danh sách EventOut."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    h = HEAD.get(locale, HEAD["vi"])
    states = STATE_LABEL.get(locale, STATE_LABEL["vi"])
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = h["sheet_rates"]
    bold = Font(bold=True)
    fills = {
        "sold_out": PatternFill("solid", fgColor="FDE2E1"),
        "restricted": PatternFill("solid", fgColor="FEF3C7"),
    }
    header = [
        h[k]
        for k in (
            "hotel",
            "role",
            "night",
            "state",
            "price",
            "rooms_left",
            "min_stay",
            "promos",
            "observed",
        )
    ]
    ws.append(header)
    for c in ws[1]:
        c.font = bold
    for row in overview.hotels:
        name = row.label or row.hotel.name or f"#{row.hotel.id}"
        for cell in row.cells:
            promos = ", ".join(f"{k} −{v}%" if v else k for k, v in (cell.promos or {}).items())
            ws.append(
                [
                    name,
                    row.role,
                    cell.stay_date,
                    states.get(cell.state, cell.state or ""),
                    _num(cell.min_price),
                    cell.exact_rooms_left,
                    cell.min_stay if cell.min_stay and cell.min_stay > 1 else None,
                    promos or None,
                    cell.last_observed_at.replace(tzinfo=None) if cell.last_observed_at else None,
                ]
            )
            fill = fills.get(cell.state or "")
            if fill:
                for c in ws[ws.max_row]:
                    c.fill = fill
    cs = wb.create_sheet(h["sheet_compset"])
    cs.append(
        [
            h[k]
            for k in (
                "night",
                "n_priced",
                "n_total",
                "sold_out",
                "restricted",
                "median",
                "own",
                "index",
                "rank",
                "sample",
            )
        ]
    )
    for c in cs[1]:
        c.font = bold
    for d in overview.compset:
        cs.append(
            [
                d.stay_date,
                d.competitors_priced,
                d.competitors_total,
                d.competitors_sold_out,
                d.competitors_restricted,
                _num(d.median_price),
                _num(d.own_min_price),
                _num(d.price_index),
                f"{d.own_rank}/{d.priced_hotels}" if d.own_rank else None,
                d.sample,
            ]
        )
    ch = wb.create_sheet(h["sheet_changes"])
    ch.append(
        [h[k] for k in ("observed", "hotel", "channel", "night", "event", "change", "reason")]
    )
    for c in ch[1]:
        c.font = bold
    for e in events:
        ch.append(
            [
                e.observed_at.replace(tzinfo=None),
                e.hotel_name,
                e.channel,
                e.stay_date,
                e.event_type,
                f"{e.from_value or ''} → {e.to_value or ''}",
                e.reason,
            ]
        )
    for sheet in (ws, cs, ch):
        for i, col in enumerate(sheet.columns, start=1):
            width = max(len(str(c.value or "")) for c in col)
            sheet.column_dimensions[get_column_letter(i)].width = min(48, max(10, width + 2))
        sheet.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@dataclass(frozen=True)
class MonthlyReport:
    tenant_name: str
    hotel_name: str
    month: date  # ngày đầu tháng
    kpis: Any  # app.pms.otb.Kpis (PMS thật)
    kpis_prev: Any | None
    index_median: Decimal | None
    index_nights: int
    outcomes: Any | None  # OutcomesOut
    highlights: Iterable[str]


def _fmt(v: Any, suffix: str = "") -> str:
    if v is None:
        return "—"
    if isinstance(v, Decimal):
        v = f"{v:,.0f}".replace(",", ".") if v >= 1000 else f"{v}"
    return f"{v}{suffix}"


def monthly_report_html(r: MonthlyReport, locale: str = "vi") -> str:
    vi = locale != "en"
    t = (
        {
            "title": "Báo cáo doanh thu phòng tháng",
            "kpi": "Chỉ số thật (PMS)",
            "occ": "Công suất",
            "adr": "ADR",
            "revpar": "RevPAR",
            "rev": "Doanh thu phòng",
            "prev": "Tháng trước",
            "pos": "Vị trí giá so compset",
            "index": "Chỉ số giá niêm yết (trung vị các đêm đủ mẫu)",
            "nights": "đêm đủ mẫu",
            "dec": "Quyết định giá",
            "applied": "Gợi ý đã áp dụng",
            "good": "đúng hướng",
            "review": "cần xem lại",
            "high": "Diễn biến đối thủ",
            "note": "Giá đối thủ là giá niêm yết công khai trên OTA (không phải ADR). "
            "ADR/RevPAR chỉ tính từ số PMS của bạn.",
            "print": "In / lưu PDF",
        }
        if vi
        else {
            "title": "Monthly rooms revenue report",
            "kpi": "Actual KPIs (PMS)",
            "occ": "Occupancy",
            "adr": "ADR",
            "revpar": "RevPAR",
            "rev": "Rooms revenue",
            "prev": "Previous month",
            "pos": "Rate position vs compset",
            "index": "Advertised rate index (median of nights with enough sample)",
            "nights": "nights with sample",
            "dec": "Pricing decisions",
            "applied": "Suggestions applied",
            "good": "right direction",
            "review": "to review",
            "high": "Competitor highlights",
            "note": "Competitor prices are public advertised rates on OTAs (not ADR). "
            "ADR/RevPAR come from your PMS only.",
            "print": "Print / save PDF",
        }
    )
    k, p = r.kpis, r.kpis_prev
    rows = [
        (t["occ"], _fmt(k.occupancy_pct, "%"), _fmt(p.occupancy_pct, "%") if p else "—"),
        (t["adr"], _fmt(k.adr), _fmt(p.adr) if p else "—"),
        (t["revpar"], _fmt(k.revpar), _fmt(p.revpar) if p else "—"),
        (t["rev"], _fmt(k.revenue), _fmt(p.revenue) if p else "—"),
    ]
    kpi_html = "".join(
        f"<tr><td>{escape(a)}</td><td><b>{escape(b)}</b></td><td>{escape(c)}</td></tr>"
        for a, b, c in rows
    )
    o = r.outcomes
    dec = (
        f"<p>{escape(t['applied'])}: <b>{o.applied}/{o.decided}</b> · "
        f"{o.good} {escape(t['good'])} · {o.review} {escape(t['review'])}</p>"
        if o is not None and o.decided
        else "<p>—</p>"
    )
    hl = "".join(f"<li>{escape(x)}</li>" for x in r.highlights) or "<li>—</li>"
    month = r.month.strftime("%m/%Y")
    return f"""<!doctype html><html lang="{"vi" if vi else "en"}"><head><meta charset="utf-8">
<title>{escape(t["title"])} {month} — {escape(r.hotel_name)}</title>
<style>body{{font:14px/1.5 -apple-system,Segoe UI,Roboto,Arial,sans-serif;color:#111827;max-width:820px;margin:32px auto;padding:0 16px}}
h1{{font-size:22px;margin:0}}h2{{font-size:16px;margin-top:28px;border-bottom:1px solid #e5e7eb;padding-bottom:4px}}
table{{border-collapse:collapse;width:100%}}td{{padding:6px 8px;border-bottom:1px solid #f1f5f9}}
.muted{{color:#6b7280;font-size:12px}}@media print{{.noprint{{display:none}}}}</style></head><body>
<p class="noprint"><button onclick="window.print()">{escape(t["print"])}</button></p>
<h1>{escape(t["title"])} {month}</h1><p class="muted">{escape(r.hotel_name)} · {escape(r.tenant_name)}</p>
<h2>{escape(t["kpi"])}</h2><table><tr><td></td><td>{month}</td><td>{escape(t["prev"])}</td></tr>{kpi_html}</table>
<h2>{escape(t["pos"])}</h2><p>{escape(t["index"])}: <b>{_fmt(r.index_median)}</b> ({r.index_nights} {escape(t["nights"])})</p>
<h2>{escape(t["dec"])}</h2>{dec}
<h2>{escape(t["high"])}</h2><ul>{hl}</ul>
<p class="muted">{escape(t["note"])}</p></body></html>"""
