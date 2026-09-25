import csv
import io
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any

from app.pms.base import (
    CANONICAL_COLUMNS,
    REQUIRED_COLUMNS,
    OwnHotelDailyRow,
    PmsAdapterError,
    RowError,
    Table,
)

# Tên cột thường gặp trong file xuất từ ezCloud, Newway, Hotel Link, Smile và Excel tự làm.
_ALIASES: dict[str, tuple[str, ...]] = {
    "stay_date": ("stay_date", "date", "ngay", "ngày", "business date", "stay date", "ngày ở"),
    "rooms_total": ("rooms_total", "total rooms", "tong phong", "tổng phòng", "capacity", "rooms"),
    "rooms_sold": (
        "rooms_sold",
        "sold",
        "rooms sold",
        "occupied",
        "phong ban",
        "phòng bán",
        "phòng đã bán",
        "room nights",
    ),
    "rooms_available": (
        "rooms_available",
        "available",
        "rooms available",
        "phong trong",
        "phòng trống",
        "vacant",
    ),
    "occupancy_pct": (
        "occupancy_pct",
        "occupancy",
        "occ",
        "occ %",
        "occupancy %",
        "cong suat",
        "công suất",
    ),
    "adr": ("adr", "average daily rate", "gia trung binh", "giá trung bình", "avg rate"),
    "revenue": ("revenue", "room revenue", "doanh thu", "doanh thu phòng"),
}

_DATE_FORMATS = ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d", "%d.%m.%Y", "%m/%d/%Y")


def _get(raw: dict[str, Any], mapping: dict[str, str], col: str) -> Any:
    src = mapping.get(col)
    return raw.get(src) if src else None


def _norm(name: str) -> str:
    return re.sub(r"\s+", " ", name.strip().lower().replace("_", " "))


def template_csv() -> str:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(CANONICAL_COLUMNS)
    w.writerow(["2026-10-01", "120", "96", "24", "80.0", "1850000", "177600000"])
    w.writerow(["2026-10-02", "120", "108", "12", "90.0", "1990000", "214920000"])
    return buf.getvalue()


def parse_date(value: Any) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"unrecognised date {text!r}")


_CURRENCY_RE = re.compile(r"^(₫|vnđ|vnd|đồng|đ)|(₫|vnđ|vnd|đồng|đ)$", re.IGNORECASE)


def _is_thousands(parts: list[str]) -> bool:
    """`parts` là các nhóm tách theo một dấu: nhóm đầu 1–3 số không bắt đầu bằng 0, các nhóm sau
    đúng 3 số ("0.500" không phải 500)."""
    head = parts[0]
    return (
        1 <= len(head) <= 3
        and head.isdigit()
        and head[0] != "0"
        and all(len(p) == 3 and p.isdigit() for p in parts[1:])
    )


def _normalize_number(text: str, money: bool) -> str:
    """Chuẩn hoá chuỗi số kiểu Việt Nam/Anh về dạng Decimal hiểu được.

    "1.850.000" và "1,850,000" là ngăn nghìn; "1.234.567,89" / "1,234,567.89" có phần thập phân;
    "850.000" (một dấu chấm, đúng 3 số sau) chỉ coi là ngăn nghìn ở cột tiền (money=True).
    Bỏ khoảng trắng, "%" và đơn vị tiền (₫, đ, VND, VNĐ, đồng) ở đầu/cuối.
    """
    text = text.replace(" ", "").replace("\u00a0", "").replace("%", "")
    text = _CURRENCY_RE.sub("", text)
    sign = "-" if text.startswith("-") else ""
    text = text.lstrip("+-")
    if "," in text and "." in text:
        if text.rfind(".") > text.rfind(","):
            return sign + text.replace(",", "")
        return sign + text.replace(".", "").replace(",", ".")
    for sep, decimal_sep in ((",", "."), (".", None)):
        if sep not in text:
            continue
        parts = text.split(sep)
        if _is_thousands(parts) and (len(parts) > 2 or sep == "," or money):
            return sign + text.replace(sep, "")
        return sign + (text.replace(sep, decimal_sep) if decimal_sep else text)
    return sign + text


def _finite(d: Decimal, value: Any, what: str) -> Decimal:
    if not d.is_finite():  # "NaN"/"Infinity" hợp lệ với Decimal nhưng không phải số liệu
        raise ValueError(f"{what} {value!r}")
    return d


def parse_int(value: Any) -> int | None:
    if value is None or str(value).strip() == "":
        return None
    text = (
        str(value)
        if isinstance(value, int | float | Decimal)
        else _normalize_number(str(value).strip(), money=True)
    )
    try:
        d = _finite(Decimal(text), value, "not an integer")
    except InvalidOperation as exc:
        raise ValueError(f"not an integer {value!r}") from exc
    if d != d.to_integral_value():
        raise ValueError(f"not an integer {value!r}")
    return int(d)


def parse_decimal(value: Any, money: bool = False) -> Decimal | None:
    if value is None or str(value).strip() == "":
        return None
    text = (
        str(value)  # ô số của Excel: giữ nguyên giá trị
        if isinstance(value, int | float | Decimal)
        else _normalize_number(str(value).strip(), money)
    )
    try:
        return _finite(Decimal(text), value, "not a number")
    except InvalidOperation as exc:
        raise ValueError(f"not a number {value!r}") from exc


def _parse_money(value: Any) -> Decimal | None:
    return parse_decimal(value, money=True)


class CsvAdapter:
    """Đọc CSV hoặc Excel theo template, có ánh xạ cột."""

    name = "csv"

    def read_table(self, content: bytes, filename: str) -> Table:
        lower = filename.lower()
        if lower.endswith((".xlsx", ".xlsm")):
            return self._read_excel(content)
        if lower.endswith(".xls") or content[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
            raise PmsAdapterError(
                "định dạng Excel cũ (.xls) không được hỗ trợ; lưu lại thành .xlsx hoặc .csv"
            )
        if content[:2] == b"PK":
            return self._read_excel(content)
        return self._read_csv(content)

    @staticmethod
    def _read_csv(content: bytes) -> Table:
        for enc in ("utf-8-sig", "utf-16", "cp1258", "latin-1"):
            try:
                text = content.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise PmsAdapterError("cannot decode file")
        sample = text[:4096]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        reader = csv.DictReader(io.StringIO(text), dialect=dialect)
        columns = [c.strip() for c in (reader.fieldnames or []) if c is not None]
        if not columns:
            raise PmsAdapterError("empty file or missing header row")
        rows = [
            {
                k.strip(): (v.strip() if isinstance(v, str) else v)
                for k, v in r.items()
                if k is not None
            }
            for r in reader
        ]
        return Table(columns=columns, rows=rows)

    @staticmethod
    def _read_excel(content: bytes) -> Table:
        from openpyxl import load_workbook

        try:
            wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:  # noqa: BLE001
            raise PmsAdapterError(f"cannot read excel: {exc}") from exc
        ws = wb.worksheets[0]
        it = ws.iter_rows(values_only=True)
        header = next(it, None)
        if not header:
            raise PmsAdapterError("empty sheet")
        columns = [str(c).strip() if c is not None else f"col{i}" for i, c in enumerate(header)]
        rows = []
        for raw in it:
            if raw is None or all(v is None for v in raw):
                continue
            rows.append({columns[i]: raw[i] for i in range(min(len(columns), len(raw)))})
        return Table(columns=columns, rows=rows)

    def suggest_mapping(self, columns: list[str]) -> dict[str, str]:
        normalised = {_norm(c): c for c in columns}
        mapping: dict[str, str] = {}
        for canonical, aliases in _ALIASES.items():
            for alias in aliases:
                if _norm(alias) in normalised:
                    mapping[canonical] = normalised[_norm(alias)]
                    break
        return mapping

    def parse(
        self, table: Table, mapping: dict[str, str]
    ) -> tuple[list[OwnHotelDailyRow], list[RowError]]:
        errors: list[RowError] = []
        for required in REQUIRED_COLUMNS:
            if required not in mapping or mapping[required] not in table.columns:
                errors.append(RowError(0, required, f"missing mapping for {required}"))
        if errors:
            return [], errors
        rows: list[OwnHotelDailyRow] = []
        seen: set[date] = set()
        for i, raw in enumerate(table.rows, start=1):
            row_errors: list[RowError] = []

            try:
                stay = parse_date(_get(raw, mapping, "stay_date"))
            except ValueError as exc:
                errors.append(RowError(i, "stay_date", str(exc)))
                continue
            if stay in seen:
                errors.append(RowError(i, "stay_date", f"duplicate date {stay.isoformat()}"))
                continue
            vals: dict[str, Any] = {}
            for col, fn in (
                ("rooms_total", parse_int),
                ("rooms_sold", parse_int),
                ("rooms_available", parse_int),
                ("occupancy_pct", parse_decimal),
                ("adr", _parse_money),
                ("revenue", _parse_money),
            ):
                try:
                    vals[col] = fn(_get(raw, mapping, col))
                except (ValueError, InvalidOperation) as exc:
                    row_errors.append(RowError(i, col, str(exc)))
                    vals[col] = None
            if row_errors:
                errors.extend(row_errors)
                continue
            total, sold, avail = vals["rooms_total"], vals["rooms_sold"], vals["rooms_available"]
            if avail is None and total is not None and sold is not None:
                avail = total - sold
            if sold is None and total is not None and avail is not None:
                sold = total - avail
            if total is not None and sold is not None and sold > total:
                errors.append(RowError(i, "rooms_sold", f"sold {sold} > total {total}"))
                continue
            occ = vals["occupancy_pct"]
            if occ is None and total and sold is not None:
                occ = (Decimal(sold) / Decimal(total) * 100).quantize(Decimal("0.01"))
            if occ is not None and not (0 <= occ <= 100):
                errors.append(RowError(i, "occupancy_pct", f"out of range {occ}"))
                continue
            seen.add(stay)
            rows.append(
                OwnHotelDailyRow(
                    stay_date=stay,
                    rooms_total=total,
                    rooms_sold=sold,
                    rooms_available=avail,
                    occupancy_pct=occ,
                    adr=vals["adr"],
                    revenue=vals["revenue"],
                )
            )
        return rows, errors
