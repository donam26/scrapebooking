"""Luồng đầu-cuối các API của roadmap chuẩn hoá RM (Phase 0–7) trên DB test."""

import io
from datetime import UTC, date, datetime, timedelta
from decimal import Decimal
from typing import Any
from zoneinfo import ZoneInfo

from httpx import AsyncClient
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.api.auth import hash_password
from app.config import Settings
from app.db.models import (
    AvailabilityEvent,
    Hotel,
    HotelDateMetric,
    Listing,
    Notification,
    NotificationDelivery,
    NotificationRecipient,
    Probe,
    ScanRun,
    Tenant,
    TenantHotel,
    User,
)
from app.notify.channels import Message, SendResult
from app.notify.service import NotificationService
from tests.fakes import FakeEmailSender
from tests.integration.test_api import _login

NOW = datetime.now(tz=UTC)
TODAY = NOW.astimezone(ZoneInfo("Asia/Ho_Chi_Minh")).date()
N1 = TODAY + timedelta(days=4)
N2 = TODAY + timedelta(days=5)


async def _seed(db: AsyncSession) -> dict[str, int]:
    t = Tenant(
        name="Rex Group",
        timezone="Asia/Ho_Chi_Minh",
        scan_times=["06:00", "14:00", "22:00"],
        horizon_days=30,
        insight_language="vi",
        insight_hour="07:30",
        country_code="vn",
        active=True,
    )
    db.add(t)
    await db.flush()
    hotels = {}
    for key, name in (("own", "Rex"), ("c1", "C1"), ("c2", "C2"), ("c3", "C3"), ("c4", "C4")):
        h = Hotel(name=name, country_code="vn")
        db.add(h)
        await db.flush()
        hotels[key] = h.id
        db.add(
            TenantHotel(
                tenant_id=t.id,
                hotel_id=h.id,
                role="self" if key == "own" else "competitor",
                active=True,
            )
        )
        db.add(
            Listing(
                hotel_id=h.id,
                channel="booking",
                listing_key=f"vn/{name.lower()}",
                url=f"https://www.booking.com/hotel/vn/{name.lower()}.html",
                status="active",
            )
        )
    db.add_all(
        [
            User(
                tenant_id=t.id,
                email="admin@rex.vn",
                password_hash=hash_password("admin-pass-1"),
                role="tenant_admin",
                active=True,
            ),
            User(
                tenant_id=t.id,
                email="view@rex.vn",
                password_hash=hash_password("view-pass-1"),
                role="viewer",
                active=True,
            ),
        ]
    )
    run = ScanRun(
        trigger_key="seed:booking",
        channel="booking",
        scheduled_at=NOW - timedelta(hours=1),
        started_at=NOW - timedelta(hours=1),
        finished_at=NOW - timedelta(minutes=30),
        status="completed",
        total_probes=10,
    )
    db.add(run)
    await db.flush()
    await db.commit()
    return {"t": t.id, "run": run.id, **hotels}


def _metric(
    hid: int, d: date, run: int, status: str, price: str | None, **kw: Any
) -> HotelDateMetric:
    return HotelDateMetric(
        hotel_id=hid,
        channel="booking",
        stay_date=d,
        as_of_scan_run_id=run,
        days_to_arrival=(d - TODAY).days,
        min_price=Decimal(price) if price else None,
        currency="VND",
        availability_status=status,
        last_observed_at=NOW - timedelta(hours=1),
        **kw,
    )


async def _market(db: AsyncSession, ids: dict[str, int]) -> None:
    run = ids["run"]
    db.add_all(
        [
            _metric(ids["own"], N1, run, "available", "900000", exact_rooms_left=6,
                    prices_by_key={"t|f": "900000"}),
            _metric(ids["c1"], N1, run, "sold_out", None),
            _metric(ids["c2"], N1, run, "available", "1000000", exact_rooms_left=2,
                    prices_by_key={"t|f": "1200000", "f|f": "1000000"},
                    promos={"Late Escape Deal": "40.0"}),
            _metric(ids["c3"], N1, run, "available", "1100000", min_stay=2,
                    prices_by_key={"t|t": "1100000"}),
            _metric(ids["c4"], N1, run, "restricted", None),
            _metric(ids["own"], N2, run, "available", "950000"),
            *(
                _metric(ids[k], N2, run, "available", p, prices_by_key={"t|f": p})
                for k, p in (("c1", "1000000"), ("c2", "1100000"), ("c3", "1200000"), ("c4", "1300000"))
            ),
        ]
    )  # fmt: skip
    for hid in (ids["own"], ids["c2"]):
        db.add(
            Probe(
                scan_run_id=run,
                hotel_id=hid,
                channel="booking",
                stay_date=N1,
                checkin=N1,
                checkout=N1 + timedelta(days=1),
                nights=1,
                adults=2,
                status="ok",
                fetched_at=NOW - timedelta(hours=1),
            )
        )
    db.add(
        AvailabilityEvent(
            hotel_id=ids["c2"],
            channel="booking",
            stay_date=N1,
            event_type="promo_start",
            to_value="Late Escape Deal",
            delta=Decimal(40),
            confidence="exact",
            scan_run_id=run,
            observed_at=NOW - timedelta(hours=1),
        )
    )
    await db.commit()


async def test_compset_sample_states_data_status_and_radar(
    client: AsyncClient, db: AsyncSession
) -> None:
    ids = await _seed(db)
    await _market(db, ids)
    await _login(client, "view@rex.vn", "view-pass-1")
    p = {"start": N1.isoformat(), "end": N2.isoformat()}
    ov = (await client.get("/overview", params=p)).json()
    c1, c2 = ov["compset"]
    # Đêm 1: 4 đối thủ, 1 hết, 1 hạn chế (lịch), 1 chỉ bán ≥2 đêm, chỉ C2 có giá 1 đêm.
    assert (c1["competitors_sold_out"], c1["competitors_restricted"]) == (1, 2)
    assert (c1["competitors_priced"], c1["sample"], c1["median_price"]) == (1, "insufficient", None)
    # Đêm 2: đủ 4 đối thủ có giá → trung vị 1.150.000, chỉ số 82,6, vị trí 1/5.
    assert (c2["sample"], c2["median_price"], c2["price_index"], c2["own_rank"]) == (
        "ok", "1150000.00", "82.6", 1,
    )  # fmt: skip
    cells = {h["hotel"]["name"]: h["cells"][0] for h in ov["hotels"]}
    assert cells["C3"]["state"] == "restricted" and cells["C3"]["min_stay"] == 2
    assert cells["C4"]["state"] == "restricted" and cells["C1"]["state"] == "sold_out"
    assert cells["C2"]["promos"] == {"Late Escape Deal": "40.0"}
    assert ov["last_run"]["id"] == ids["run"]

    ds = (await client.get("/data-status")).json()
    assert ds["stale"] is False and ds["stale_after_hours"] == 10
    assert ds["success_rate_7d"] == "1.0" and ds["probes_7d"] == 2

    promos = (await client.get("/market/radar/promotions", params=p)).json()
    c2promo = next(h for h in promos["hotels"] if h["name"] == "C2")["promos"][0]
    assert c2promo["label"] == "Late Escape Deal" and c2promo["origin"] == "hotel"
    assert c2promo["max_depth_pct"] == "40.0" and c2promo["started_at"] is not None
    restr = (await client.get("/market/radar/restrictions", params=p)).json()
    by = {h["name"]: h["nights"] for h in restr["hotels"]}
    assert by["C3"][0]["min_stay"] == 2 and by["C4"][0]["closed_to_arrival"] is True
    assert "closed_channels" not in by["C3"][0]
    canc = (await client.get("/market/radar/cancellation", params=p)).json()
    c2c = next(h for h in canc["hotels"] if h["name"] == "C2")
    assert c2c["nr_discount_pct"] == "16.7" and c2c["refundable_share"] == "1.00"
    assert (await client.get("/market/radar/area-scarcity")).status_code == 404  # chưa có khu vực

    # Excel rate shop: 3 trang, có trạng thái "Hạn chế".
    x = await client.get("/export/rate-shop.xlsx", params=p)
    assert x.status_code == 200, x.text
    wb = load_workbook(io.BytesIO(x.content))
    assert wb.sheetnames == ["Giá đối thủ", "Compset", "Thay đổi"]
    states = {r[3] for r in wb["Giá đối thủ"].iter_rows(min_row=2, values_only=True)}
    assert {"Hạn chế", "Hết phòng", "Còn bán"} <= states


async def test_watchlist_tiers_rooms_total_and_compset_review(
    client: AsyncClient, db: AsyncSession
) -> None:
    ids = await _seed(db)
    await _login(client, "admin@rex.vn", "admin-pass-1")
    r = await client.patch(
        f"/watchlist/{ids['c4']}", json={"tier": "secondary", "rooms_total": 400}
    )
    assert r.status_code == 200 and r.json()["tier"] == "secondary"
    assert r.json()["hotel"]["rooms_total"] == 400
    bad = await client.patch(f"/watchlist/{ids['c1']}", json={"compset_of": ids["c2"]})
    assert bad.status_code == 422
    review = (await client.get("/watchlist/compset-review")).json()
    assert review["primary"] == 3 and review["secondary"] == 1
    assert review["warnings"] == ["too_few", "rooms_unknown"]
    s = await client.patch("/settings", json={"source_markets": ["KR", "cn"]})
    assert s.status_code == 200 and s.json()["source_markets"] == ["cn", "kr"]
    assert (await client.patch("/settings", json={"source_markets": ["xx"]})).status_code == 422


async def test_otb_import_pickup_pace_and_kpis(client: AsyncClient, db: AsyncSession) -> None:
    ids = await _seed(db)
    await _login(client, "admin@rex.vn", "admin-pass-1")
    stay = TODAY + timedelta(days=10)
    csv = (
        "Ngày đặt,Ngày đến,Ngày đi,Số phòng,Doanh thu,Trạng thái\n"
        f"{TODAY - timedelta(days=20)},{stay},{stay + timedelta(days=2)},2,4000000,confirmed\n"
        f"{TODAY - timedelta(days=3)},{stay},{stay + timedelta(days=1)},1,1500000,confirmed\n"
        f"{TODAY - timedelta(days=10)},{stay},{stay + timedelta(days=1)},5,6000000,cancelled\n"
        f"{TODAY},{stay},{stay + timedelta(days=1)},1,1600000,confirmed\n"
    )
    r = await client.post(
        "/pms/otb/import",
        data={"hotel_id": str(ids["own"]), "kind": "bookings", "rooms_available": "20"},
        files={"file": ("bookings.csv", csv.encode(), "text/csv")},
    )
    assert r.status_code == 201, r.text
    assert r.json()["status"] == "completed" and r.json()["snapshots_written"] > 10
    competitor = await client.post(
        "/pms/otb/import",
        data={"hotel_id": str(ids["c1"]), "kind": "otb_report"},
        files={"file": ("x.csv", b"stay_date,rooms_otb\n2026-10-17,3\n", "text/csv")},
    )
    assert competitor.status_code == 422
    body = (
        await client.get("/market/otb", params={"start": stay.isoformat(), "end": stay.isoformat()})
    ).json()
    n = body["nights"][0]
    # Hôm nay: 2 + 1 + 1 = 4 phòng (đặt 5 đã huỷ không tính); hôm qua 3; 7 ngày trước 2.
    assert (n["rooms_otb"], n["pickup_1d"], n["pickup_7d"]) == (4, 1, 2)
    assert n["occ_otb_pct"] == "20.0" and n["revenue_otb"] == "5100000.00"
    assert body["as_of_date"] == TODAY.isoformat()


async def test_subscriptions_zalo_fanout_tracking_and_engagement(
    client: AsyncClient, db: AsyncSession, settings: Settings
) -> None:
    ids = await _seed(db)
    db.add(NotificationRecipient(tenant_id=ids["t"], email="gm@rex.vn", active=True))
    await db.commit()
    await _login(client, "view@rex.vn", "view-pass-1")
    sub = await client.post(
        "/notifications/subscriptions",
        json={"channel": "zalo", "target": "0912 345 678", "kinds": ["alerts", "test"]},
    )
    assert sub.status_code == 422  # "test" không phải loại đăng ký được
    sub = await client.post(
        "/notifications/subscriptions",
        json={"channel": "zalo", "target": "0912 345 678", "kinds": ["alerts"], "max_per_day": 5},
    )
    assert sub.status_code == 201, sub.text
    assert sub.json()["target"] == "84912345678"
    bad = await client.post(
        "/notifications/subscriptions", json={"channel": "webhook", "target": "http://x"}
    )
    assert bad.status_code == 422
    assert len((await client.get("/notifications/subscriptions")).json()) == 1

    class FakeZalo:
        channel = "zalo"
        configured = True
        sent: list[tuple[str, Message, str]] = []

        async def send(self, target: str, message: Message, tracking_id: str) -> SendResult:
            self.sent.append((target, message, tracking_id))
            return SendResult("msg-1", Decimal(300))

    zalo = FakeZalo()
    svc = NotificationService(db, FakeEmailSender(), settings, notifiers={"zalo": zalo})
    email = await svc._send(  # gửi một tin cảnh báo như dispatch_alerts
        ids["t"],
        "alerts",
        f"alerts:{ids['t']}:{ids['run']}",
        __import__("app.notify.render", fromlist=["Email"]).Email(
            "Đêm T7 đang căng", "Đêm T7 đang căng\n\n• 3/4 đối thủ hết", '<a href="x">x</a>'
        ),
        NOW,
        item_count=1,
        scan_run_id=ids["run"],
    )
    await db.commit()
    assert email is not None and email.status == "sent" and email.channel == "email"
    rows = (await db.execute(select(Notification).order_by(Notification.id))).scalars().all()
    assert [(r.channel, r.status) for r in rows] == [("email", "sent"), ("zalo", "sent")]
    assert zalo.sent[0][0] == "84912345678" and "/api/notifications/t/" in zalo.sent[0][1].url
    delivery = (
        await db.execute(select(NotificationDelivery).where(NotificationDelivery.channel == "zalo"))
    ).scalar_one()
    assert delivery.cost_vnd == Decimal(300) and delivery.external_id == "msg-1"

    click = await client.get(f"/notifications/t/{delivery.token}", params={"to": "/today"})
    assert click.status_code == 302 and click.headers["location"] == "/today"
    evil = await client.get(f"/notifications/t/{delivery.token}", params={"to": "//evil.com"})
    assert evil.headers["location"] == "/today"
    done = await client.get(f"/notifications/t/{delivery.token}/resolve")
    assert done.status_code == 302
    eng = (await client.get("/notifications/engagement")).json()["weeks"]
    z = next(w for w in eng if w["channel"] == "zalo")
    assert (z["sent"], z["clicked"], z["resolved"], z["cost_vnd"]) == (1, 1, 1, "300.00")
